from pptx import Presentation
ppt=r'C:\Users\zjuat\Desktop\无视频PPT.pptx'; out=r'C:\Users\zjuat\Desktop\ppt_text.txt'
prs=Presentation(ppt)
with open(out,'w',encoding='utf-8') as f:
 for i,slide in enumerate(prs.slides,1):
  f.write(f'\n### 第{i}页\n')
  for sh in slide.shapes:
   if hasattr(sh,'text') and sh.text.strip(): f.write(sh.text.replace('\n',' | ')+'\n')
print(out)
